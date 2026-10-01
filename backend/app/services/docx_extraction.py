from io import BytesIO
from zipfile import BadZipFile, ZipFile, is_zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.contracts.attachment_content import HashedAttachmentContent
from app.contracts.attachment_extraction import AttachmentExtractionResult, DOCXTextSegment, ExtractionMethod, ExtractionReason
from app.models.enums import AttachmentProcessingState

DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def extract_docx(
    content: HashedAttachmentContent,
    *,
    filename: str,
    mime_type: str,
    max_paragraphs: int,
    max_tables: int,
    max_table_cells: int,
    max_characters: int,
) -> AttachmentExtractionResult:
    if not content.content:
        return _failed(content, ExtractionReason.EMPTY_ATTACHMENT)
    if not filename.casefold().endswith(".docx") or mime_type.casefold() != DOCX_MIME_TYPE:
        return _failed(content, ExtractionReason.CONTENT_TYPE_MISMATCH)
    if min(max_paragraphs, max_tables, max_table_cells, max_characters) <= 0:
        raise ValueError("DOCX extraction limits must be positive")
    if not _is_docx_package(content.content):
        return _failed(content, ExtractionReason.CONTENT_TYPE_MISMATCH)

    try:
        document = Document(BytesIO(content.content))
    except (BadZipFile, KeyError, PackageNotFoundError, SyntaxError, ValueError):
        return _failed(content, ExtractionReason.MALFORMED_DOCX)

    parts: list[str] = []
    segments: list[DOCXTextSegment] = []
    text_length = 0
    processed_units = 0
    truncated = len(document.paragraphs) > max_paragraphs or len(document.tables) > max_tables

    def add_segment(text: str, **coordinates: int) -> bool:
        nonlocal text_length, truncated
        normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not normalized:
            return True
        separator = "\n\n" if text_length else ""
        remaining = max_characters - text_length
        if len(separator) + len(normalized) > remaining:
            truncated = True
            normalized = normalized[: max(0, remaining - len(separator))]
            separator = separator if normalized else ""
        if not normalized:
            return False
        parts.append(separator + normalized)
        start = text_length + len(separator)
        text_length += len(separator) + len(normalized)
        segments.append(
            DOCXTextSegment(
                text_start=start,
                text_end=text_length,
                text=normalized,
                **coordinates,
            )
        )
        return text_length < max_characters

    keep_processing = True
    for paragraph_index, paragraph in enumerate(document.paragraphs[:max_paragraphs]):
        processed_units += 1
        keep_processing = add_segment(paragraph.text, paragraph_index=paragraph_index)
        if not keep_processing:
            break

    processed_cells = 0
    if keep_processing:
        for table_index, table in enumerate(document.tables[:max_tables]):
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    if processed_cells == max_table_cells:
                        truncated = True
                        keep_processing = False
                        break
                    processed_cells += 1
                    processed_units += 1
                    keep_processing = add_segment(
                        cell.text,
                        table_index=table_index,
                        row_index=row_index,
                        column_index=column_index,
                    )
                    if not keep_processing:
                        break
                if not keep_processing:
                    break
            if not keep_processing:
                break

    extracted_text = "".join(parts)
    meaningful = any(character.isalnum() for character in extracted_text)
    return AttachmentExtractionResult(
        attachment_id=content.attachment_id,
        content_hash=content.content_hash,
        status=(AttachmentProcessingState.EXTRACTED if meaningful else AttachmentProcessingState.FAILED),
        extraction_method=ExtractionMethod.PYTHON_DOCX,
        extracted_text=extracted_text if meaningful else None,
        docx_segments=tuple(segments) if meaningful else (),
        reason=(
            ExtractionReason.EXTRACTION_LIMIT_REACHED
            if meaningful and truncated
            else None if meaningful
            else ExtractionReason.EMPTY_DOCX
        ),
        truncated=truncated,
        source_size_bytes=content.size_bytes,
        processed_unit_count=processed_units,
    )


def _is_docx_package(data: bytes) -> bool:
    stream = BytesIO(data)
    if not is_zipfile(stream):
        return False
    try:
        with ZipFile(stream) as archive:
            names = set(archive.namelist())
    except BadZipFile:
        return False
    return {"[Content_Types].xml", "word/document.xml"} <= names


def _failed(content: HashedAttachmentContent, reason: ExtractionReason) -> AttachmentExtractionResult:
    return AttachmentExtractionResult(
        attachment_id=content.attachment_id,
        content_hash=content.content_hash,
        status=AttachmentProcessingState.FAILED,
        extraction_method=ExtractionMethod.PYTHON_DOCX,
        extracted_text=None,
        reason=reason,
        source_size_bytes=content.size_bytes,
        processed_unit_count=0,
    )
