from io import BytesIO
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from app.contracts.attachment_content import AttachmentContent
from app.contracts.attachment_extraction import ExtractionMethod, ExtractionReason
from app.models.enums import AttachmentProcessingState
from app.services.attachment_hashing import hash_attachment_content
from app.services.docx_extraction import DOCX_MIME_TYPE, extract_docx


def _docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        docx_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, text in enumerate(row):
                docx_table.cell(row_index, column_index).text = text
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _with_malformed_document_xml(data: bytes) -> bytes:
    source = BytesIO(data)
    target = BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as corrupted:
        for item in original.infolist():
            payload = b"<invalid" if item.filename == "word/document.xml" else original.read(item)
            corrupted.writestr(item, payload)
    return target.getvalue()


def _extract(data: bytes, **limits: int):
    content = hash_attachment_content(
        AttachmentContent(attachment_id=uuid4(), content=data, size_bytes=len(data))
    )
    return extract_docx(
        content,
        filename="evidence.DOCX",
        mime_type=DOCX_MIME_TYPE.upper(),
        max_paragraphs=limits.get("max_paragraphs", 10_000),
        max_tables=limits.get("max_tables", 500),
        max_table_cells=limits.get("max_table_cells", 50_000),
        max_characters=limits.get("max_characters", 1_000_000),
    )


def test_extracts_paragraphs_and_table_cells_with_provenance() -> None:
    result = _extract(_docx("First", "Second", table=[["A1", "B1"], ["A2", "B2"]]))

    assert result.status is AttachmentProcessingState.EXTRACTED
    assert result.extraction_method is ExtractionMethod.PYTHON_DOCX
    assert result.extracted_text == "First\n\nSecond\n\nA1\n\nB1\n\nA2\n\nB2"
    assert result.processed_unit_count == 6
    assert [segment.paragraph_index for segment in result.docx_segments[:2]] == [0, 1]
    assert (
        result.docx_segments[2].table_index,
        result.docx_segments[2].row_index,
        result.docx_segments[2].column_index,
    ) == (0, 0, 0)
    assert all(
        result.extracted_text[segment.text_start : segment.text_end] == segment.text
        for segment in result.docx_segments
    )


def test_empty_docx_fails_without_pdf_fallback() -> None:
    result = _extract(_docx("", "!?"))

    assert result.status is AttachmentProcessingState.FAILED
    assert result.reason is ExtractionReason.EMPTY_DOCX
    assert result.extracted_text is None
    assert result.status is not AttachmentProcessingState.FALLBACK_REQUIRED


def test_paragraph_and_character_limits_are_bounded() -> None:
    paragraph_limited = _extract(_docx("one", "two"), max_paragraphs=1)
    character_limited = _extract(_docx("abcdef"), max_characters=3)

    assert paragraph_limited.extracted_text == "one"
    assert paragraph_limited.reason is ExtractionReason.EXTRACTION_LIMIT_REACHED
    assert paragraph_limited.truncated is True
    assert character_limited.extracted_text == "abc"
    assert character_limited.reason is ExtractionReason.EXTRACTION_LIMIT_REACHED


def test_table_and_cell_limits_are_bounded() -> None:
    data = _docx(table=[["A", "B"], ["C", "D"]])
    no_tables = _extract(data, max_tables=1, max_table_cells=1)

    assert no_tables.extracted_text == "A"
    assert no_tables.processed_unit_count == 1
    assert no_tables.truncated is True


def test_rejects_type_mismatch_and_invalid_container() -> None:
    data = _docx("text")
    content = hash_attachment_content(
        AttachmentContent(attachment_id=uuid4(), content=data, size_bytes=len(data))
    )
    mismatch = extract_docx(
        content,
        filename="file.pdf",
        mime_type=DOCX_MIME_TYPE,
        max_paragraphs=1,
        max_tables=1,
        max_table_cells=1,
        max_characters=10,
    )
    invalid = _extract(b"not a DOCX")

    assert mismatch.reason is ExtractionReason.CONTENT_TYPE_MISMATCH
    assert invalid.reason is ExtractionReason.CONTENT_TYPE_MISMATCH


def test_empty_and_malformed_docx_have_distinct_reasons() -> None:
    empty = _extract(b"")
    malformed = _extract(_with_malformed_document_xml(_docx("text")))

    assert empty.reason is ExtractionReason.EMPTY_ATTACHMENT
    assert malformed.reason is ExtractionReason.MALFORMED_DOCX


@pytest.mark.parametrize("limit", ["max_paragraphs", "max_tables", "max_table_cells", "max_characters"])
def test_rejects_nonpositive_limits(limit: str) -> None:
    with pytest.raises(ValueError, match="limits must be positive"):
        _extract(_docx("text"), **{limit: 0})
